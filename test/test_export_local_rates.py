"""验证 local rate 工作簿能够确定性导出为可审计的插件费率目录。"""

from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts.export_local_rates import export_catalog


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEST_OUTPUT_ROOT = PROJECT_ROOT / "temp" / "test_vendor_local_rate"


class ExportLocalRatesTest(unittest.TestCase):
    """覆盖费率字段映射、来源追踪、缺失语义和确定性。"""

    def setUp(self) -> None:
        """创建包含四个 vendor 的最小真实结构工作簿。"""
        TEST_OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
        self.temp_dir = tempfile.TemporaryDirectory(dir=TEST_OUTPUT_ROOT)
        self.workbook_path = Path(self.temp_dir.name) / "local rate 费用表.xlsx"

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._append_sheet(
            workbook,
            "RWS",
            [
                "Source",
                "Target",
                "TEP (USD/1000 words)",
                "人工翻译\nHuman Translation rate",
                "人工校准\nPost Edit rate",
                "AI翻译+人工校准AIPE rate",
                "Hourly rate (USD/h)",
            ],
            ["English", "ro", 123, 92, 43, "74.00 ", 40],
        )
        self._append_sheet(
            workbook,
            "Pactera",
            [
                "Source",
                "Target",
                "TEP (USD/1000 words)",
                "T rate",
                "P rate",
                "AIPE rate",
                "AIPE + review rate",
                "Hourly rate (USD/h)",
            ],
            ["English", "Japanese", 100, 75, 35, 55, 70, 30],
        )
        self._append_sheet(
            workbook,
            "Alconost",
            ["Source ", "Target", "Hourly Rate/USD"],
            ["English", "Japanese (Japan)", 35],
        )
        self._append_sheet(
            workbook,
            "Skrivanek",
            ["Source ", "Target", "Hourly Rate/USD"],
            ["English", "Japanese (Japan)", 30],
        )
        workbook.save(self.workbook_path)

    def tearDown(self) -> None:
        """移除当前测试创建的临时工作簿。"""
        self.temp_dir.cleanup()

    @staticmethod
    def _append_sheet(
        workbook: Workbook,
        name: str,
        headers: list[str],
        values: list[object],
    ) -> None:
        """向 fixture 工作簿写入一个 vendor sheet。"""
        worksheet = workbook.create_sheet(name)
        worksheet.append(headers)
        worksheet.append(values)

    def test_exports_vendor_word_rates_without_derivation(self) -> None:
        """按原列导出 RWS/Pactera，并保持 hourly vendor 无按字费率。"""
        catalog = export_catalog(self.workbook_path)

        self.assertEqual(
            set(catalog["vendors"]),
            {"RWS", "Pactera", "Alconost", "Skrivanek"},
        )
        rws = catalog["vendors"]["RWS"][0]
        pactera = catalog["vendors"]["Pactera"][0]
        alconost = catalog["vendors"]["Alconost"][0]
        self.assertEqual(rws["aipe_usd_per_1000_words"], 74)
        self.assertEqual(rws["t_only_usd_per_1000_words"], 92)
        self.assertEqual(rws["source_locale"], "en")
        self.assertEqual(rws["target_locale"], "ro-RO")
        self.assertEqual(pactera["aipe_usd_per_1000_words"], 55)
        self.assertEqual(pactera["t_only_usd_per_1000_words"], 75)
        self.assertEqual(pactera["target_locale"], "ja-JP")
        self.assertNotIn("aipe_usd_per_1000_words", alconost)
        self.assertNotIn("t_only_usd_per_1000_words", alconost)
        self.assertEqual(alconost["hourly_rate_usd"], 35)

    def test_records_workbook_hash_and_source_cells(self) -> None:
        """目录保留源文件散列、sheet、行号和原始费率列名。"""
        catalog = export_catalog(self.workbook_path)
        expected_hash = hashlib.sha256(self.workbook_path.read_bytes()).hexdigest()
        rws = catalog["vendors"]["RWS"][0]

        self.assertEqual(catalog["schema_version"], 1)
        self.assertEqual(catalog["source"]["filename"], self.workbook_path.name)
        self.assertEqual(catalog["source"]["sha256"], expected_hash)
        self.assertEqual(rws["provenance"]["sheet"], "RWS")
        self.assertEqual(rws["provenance"]["row"], 2)
        self.assertEqual(
            rws["provenance"]["rate_columns"],
            {
                "aipe": "AI翻译+人工校准AIPE rate",
                "t_only": "人工翻译\nHuman Translation rate",
            },
        )

    def test_export_is_deterministic(self) -> None:
        """相同工作簿重复导出必须得到完全相同的数据对象。"""
        self.assertEqual(
            export_catalog(self.workbook_path),
            export_catalog(self.workbook_path),
        )

    def test_unknown_language_label_fails_explicitly(self) -> None:
        """工作簿新增未知语言标签时必须停止导出。"""
        workbook = Workbook()
        workbook.remove(workbook.active)
        for vendor in ("RWS", "Pactera", "Alconost", "Skrivanek"):
            if vendor == "RWS":
                headers = [
                    "Source", "Target", "人工翻译\nHuman Translation rate",
                    "AI翻译+人工校准AIPE rate", "Hourly rate (USD/h)",
                ]
                values = ["English", "Unknown Locale", 1, 1, 1]
            elif vendor == "Pactera":
                headers = ["Source", "Target", "T rate", "AIPE rate", "Hourly rate (USD/h)"]
                values = ["English", "Japanese", 1, 1, 1]
            else:
                headers = ["Source ", "Target", "Hourly Rate/USD"]
                values = ["English", "Japanese (Japan)", 1]
            self._append_sheet(workbook, vendor, headers, values)
        unknown_path = Path(self.temp_dir.name) / "unknown.xlsx"
        workbook.save(unknown_path)
        with self.assertRaisesRegex(ValueError, "unknown locale label"):
            export_catalog(unknown_path)


if __name__ == "__main__":
    unittest.main()
