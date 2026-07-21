"""把业务确认的 local rate Excel 确定性导出为插件可读取的费率目录。"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


VENDORS = ("RWS", "Pactera", "Alconost", "Skrivanek")
WORD_RATE_COLUMNS = {
    "RWS": {
        "t_only": "人工翻译\nHuman Translation rate",
        "aipe": "AI翻译+人工校准AIPE rate",
    },
    "Pactera": {"t_only": "T rate", "aipe": "AIPE rate"},
}
HOURLY_RATE_COLUMNS = {
    "RWS": "Hourly rate (USD/h)",
    "Pactera": "Hourly rate (USD/h)",
    "Alconost": "Hourly Rate/USD",
    "Skrivanek": "Hourly Rate/USD",
}

# 工作簿语言标签到 Starling API locale 的显式业务映射；未知标签必须停止导出。
LOCALE_BY_LABEL = {
    "English": "en", "English(US)": "en", "English": "en",
    "Chinese": "zh-Hans", "Arabic": "ar", "Azerbaijani": "az",
    "Bulgarian": "bg", "Bengali": "bn-IN", "Bengali (India)": "bn-IN",
    "Catalan": "ca", "Cebuano(Sinugboanon)": "ceb-PH",
    "Cebuano (Philippines)": "ceb-PH", "Czech (Czech Republic)": "cs-CZ",
    "Danish": "da", "German (Germany)": "de-DE", "Greek (Greece)": "el-GR",
    "English(GB)": "en-GB", "English (United Kingdom)": "en-GB",
    "Spanish": "es", "Spanish (Latin America)": "es-419", "Estonian": "et",
    "Farsi": "fa-IR", "Persian (Iran)": "fa-IR", "Finnish (Finland)": "fi-FI",
    "Filipino": "fil-PH", "Filipino (Philippines)": "fil-PH", "French": "fr",
    "French (Canada)": "fr-CA", "Irish": "ga", "Hebrew": "he-IL",
    "Hebrew (Israel)": "he-IL", "Hindi": "hi-IN", "Hindi (India)": "hi-IN",
    "Croatian": "hr", "Hungarian (Hungary)": "hu-HU", "Indonesian": "id-ID",
    "Indonesian (Indonesia)": "id-ID", "Icelandic": "is", "Italian (Italy)": "it-IT",
    "Japanese": "ja-JP", "Japanese (Japan)": "ja-JP", "Javanese": "jv-ID",
    "Javanese (Indonesia)": "jv-ID", "Kazakh": "kk", "Khmer": "km-KH",
    "Khmer (Cambodia)": "km-KH", "Korean": "ko-KR",
    "Korean (South Korea)": "ko-KR", "Lithuanian": "lt", "Latvian": "lv",
    "Malay/Bahasa": "ms-MY", "Malay (Malaysia)": "ms-MY", "Burmese": "my-MM",
    "Burmese (Myanmar)": "my-MM", "Norwegian Bokmål": "nb",
    "Dutch (Netherlands)": "nl-NL", "Polish (Poland)": "pl-PL",
    "Portuguese": "pt", "Portuguese (Brazil)": "pt-BR",
    "Romanian (Romania)": "ro-RO", "Russian (Russia)": "ru-RU", "Slovak": "sk",
    "Slovenian": "sl", "Albanian": "sq", "Swedish (Sweden)": "sv-SE",
    "Swahili": "sw", "Thai": "th-TH", "Thai (Thailand)": "th-TH",
    "Turkish (Turkey)": "tr-TR", "Ukrainian (Ukraine)": "uk-UA",
    "Urdu(Pak)": "ur", "Urdu": "ur", "Uzbek": "uz", "Vietnamese": "vi-VN",
    "Vietnamese (Vietnam)": "vi-VN", "Chinese S": "zh-Hans",
    "Chinese (Simplified)": "zh-Hans", "Chinese TW": "zh-Hant-TW",
    "Chinese (Traditional, Taiwan)": "zh-Hant-TW",
    "ro": "ro-RO", "pl": "pl-PL", "it": "it-IT", "sk": "sk", "tr": "tr-TR",
    "pt-BR": "pt-BR", "fr-CA": "fr-CA", "fr": "fr", "da": "da", "lv": "lv",
    "uz": "uz", "lt": "lt", "hu": "hu-HU", "es": "es", "el": "el-GR",
    "nb": "nb", "sv": "sv-SE", "kk": "kk", "ru": "ru-RU", "uk": "uk-UA",
    "nl": "nl-NL", "de": "de-DE", "et": "et", "cs": "cs-CZ", "fi": "fi-FI",
    "bg": "bg", "ca": "ca", "pt": "pt", "sq": "sq", "sl": "sl", "is": "is",
    "sw": "sw", "az": "az", "es-419": "es-419", "hr": "hr",
}


def canonical_locale(label: str, location: str) -> str:
    """把工作簿语言标签转换为 Starling locale，未知值明确失败。"""
    if label not in LOCALE_BY_LABEL:
        raise ValueError(f"{location}: unknown locale label {label!r}")
    return LOCALE_BY_LABEL[label]


def sha256_file(path: Path) -> str:
    """计算源工作簿的 SHA-256，供费率版本审计。"""
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_header(value: object) -> str:
    """去除表头两端空白，同时保留原表头内部换行。"""
    return str(value).strip() if value is not None else ""


def header_index(worksheet: Worksheet) -> dict[str, tuple[int, str]]:
    """建立规范化表头到列号和原始表头的映射，并拒绝重复表头。"""
    result: dict[str, tuple[int, str]] = {}
    for cell in worksheet[1]:
        normalized = normalize_header(cell.value)
        if not normalized:
            continue
        if normalized in result:
            raise ValueError(f"{worksheet.title}: duplicate header {normalized!r}")
        result[normalized] = (cell.column, str(cell.value))
    return result


def require_column(
    worksheet: Worksheet,
    headers: dict[str, tuple[int, str]],
    name: str,
) -> tuple[int, str]:
    """返回必需列；工作簿结构漂移时明确失败。"""
    normalized = normalize_header(name)
    if normalized not in headers:
        raise ValueError(f"{worksheet.title}: missing required column {name!r}")
    return headers[normalized]


def parse_number(value: object, location: str) -> float | int:
    """把 Excel 数值或可解析字符串转换为数值，拒绝空值和非数值。"""
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{location}: expected numeric rate, got {value!r}")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        cleaned = str(value).strip().replace(",", "").replace("$", "")
        try:
            number = float(cleaned)
        except ValueError as exc:
            raise ValueError(
                f"{location}: expected numeric rate, got {value!r}"
            ) from exc
    return int(number) if number.is_integer() else number


def export_vendor_sheet(worksheet: Worksheet, vendor: str) -> list[dict[str, Any]]:
    """按 vendor 的明确列定义导出全部非空语言对。"""
    headers = header_index(worksheet)
    source_column, _ = require_column(worksheet, headers, "Source")
    target_column, _ = require_column(worksheet, headers, "Target")
    hourly_name = HOURLY_RATE_COLUMNS[vendor]
    hourly_column, hourly_original = require_column(worksheet, headers, hourly_name)

    word_columns: dict[str, tuple[int, str]] = {}
    if vendor in WORD_RATE_COLUMNS:
        for rate_type, column_name in WORD_RATE_COLUMNS[vendor].items():
            word_columns[rate_type] = require_column(worksheet, headers, column_name)

    records: list[dict[str, Any]] = []
    for row in range(2, worksheet.max_row + 1):
        source_value = worksheet.cell(row, source_column).value
        target_value = worksheet.cell(row, target_column).value
        if source_value is None and target_value is None:
            continue
        source_language = str(source_value).strip() if source_value is not None else ""
        target_language = str(target_value).strip() if target_value is not None else ""
        if not source_language or not target_language:
            raise ValueError(
                f"{worksheet.title}!{row}: source and target language are both required"
            )

        record: dict[str, Any] = {
            "source_language": source_language,
            "target_language": target_language,
            "source_locale": canonical_locale(
                source_language,
                f"{worksheet.title}!{worksheet.cell(row, source_column).coordinate}",
            ),
            "target_locale": canonical_locale(
                target_language,
                f"{worksheet.title}!{worksheet.cell(row, target_column).coordinate}",
            ),
            "hourly_rate_usd": parse_number(
                worksheet.cell(row, hourly_column).value,
                f"{worksheet.title}!{worksheet.cell(row, hourly_column).coordinate}",
            ),
            "provenance": {
                "sheet": worksheet.title,
                "row": row,
                "language_cells": {
                    "source": worksheet.cell(row, source_column).coordinate,
                    "target": worksheet.cell(row, target_column).coordinate,
                },
                "hourly_column": hourly_original,
                "hourly_cell": worksheet.cell(row, hourly_column).coordinate,
                "rate_columns": {},
                "rate_cells": {},
            },
        }

        if word_columns:
            for rate_type, (column, original_name) in word_columns.items():
                output_key = f"{rate_type}_usd_per_1000_words"
                record[output_key] = parse_number(
                    worksheet.cell(row, column).value,
                    f"{worksheet.title}!{worksheet.cell(row, column).coordinate}",
                )
                record["provenance"]["rate_columns"][rate_type] = original_name
                record["provenance"]["rate_cells"][rate_type] = worksheet.cell(
                    row, column
                ).coordinate
        records.append(record)
    return records


def export_catalog(input_path: Path) -> dict[str, object]:
    """把已冻结的 local rate workbook 转换为无时间戳的确定性目录。"""
    path = input_path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Local rate workbook not found: {path}")
    workbook = load_workbook(path, data_only=True, read_only=False)
    missing_sheets = [vendor for vendor in VENDORS if vendor not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f"Missing vendor sheets: {', '.join(missing_sheets)}")
    vendors = {
        vendor: export_vendor_sheet(workbook[vendor], vendor) for vendor in VENDORS
    }
    return {
        "schema_version": 1,
        "source": {
            "filename": path.name,
            "sha256": sha256_file(path),
            "sheets": {vendor: len(vendors[vendor]) for vendor in VENDORS},
        },
        "vendors": vendors,
    }


def write_catalog(catalog: dict[str, object], output_path: Path) -> None:
    """以稳定缩进和键顺序写出 UTF-8 JSON。"""
    destination = output_path.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    """解析工作簿输入和 JSON 输出路径。"""
    parser = argparse.ArgumentParser(
        description="Export local rate workbook to an auditable JSON catalog."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def main() -> None:
    """执行导出并打印可供 CLI 验证的版本与行数摘要。"""
    args = parse_args()
    catalog = export_catalog(args.input)
    write_catalog(catalog, args.output)
    counts = catalog["source"]["sheets"]
    print(f"sha256={catalog['source']['sha256']}")
    print("rows=" + ", ".join(f"{vendor}:{counts[vendor]}" for vendor in VENDORS))
    print(f"output={args.output.expanduser().resolve()}")


if __name__ == "__main__":
    main()
